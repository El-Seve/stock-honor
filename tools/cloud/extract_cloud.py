#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""extract_cloud.py - extractor en Python (para la Action en la nube, sin Excel).

Replica EXACTAMENTE tools/extract.ps1 (Excel + C#): mismo filtro de alcance, mismas variantes
canonicas por mayusculas/minusculas, mismo orden de productos/PDV/stock y el mismo JSON
(ASCII-only, byte a byte). La igualdad se comprueba con tools/cloud/selftest.py contra los
resultados esperados de tests/golden/, que genera el extractor local con libros sinteticos.

Uso:
    python tools/cloud/extract_cloud.py --input archivo.xlsb --out tools/stock_data.json [--coverage]
"""
import argparse
import os
import sys

SHEET = "Base Stocks"

EXPECTED_HEADERS = [
    "PERIODO", "DIA", "UBICACIONFISICA", "ALMACEN", "CODSUBALMACEN", "CODIGOORACLE", "ITEMS",
    "VALORIZADO", "TIPOALMACENAMIENTO", "DESPLEGADO", "CANAL", "ESTADO", "MODALIDAD", "PROCESO",
    "DISPONIBILIDAD", "FAMILIA", "MARCAMODELO", "PRODUCTO", "USO", "CONDICION", "TAG1", "TAG2",
    "TIPO", "SUBTIPO", "STATUS", "BLOQUE", "MARCA", "CLASIFICACION", "GAMA", "SUBGAMA",
    "RANGOAGING", "SEMANA", "ACTIVOREPOSICION", "PUNTODEVENTA", "ESTADO_PDV",
    "CATEGORIA_EQUIPO", "CATEGORIA_ACCE", "REGION", "DEPARTAMENTO",
]

# columnas (1-based) que se usan
C_PERIODO, C_DIA, C_UBIC = 1, 2, 3
C_CODORACLE, C_ITEMS = 6, 7
C_CANAL, C_ESTADO, C_DISP = 11, 12, 15
C_FAM, C_MODELO, C_PROD, C_USO = 16, 17, 18, 19
C_TIPO, C_SUBTIPO, C_MARCA = 23, 24, 27
C_PDV, C_ESTADOPDV, C_CATEQ, C_CATACC, C_REGION, C_DEPTO = 34, 35, 36, 37, 38, 39
NEEDED = [1, 2, 3, 6, 7, 11, 12, 15, 16, 17, 18, 19, 23, 24, 27, 34, 35, 36, 37, 38, 39]
POS = dict((c, i) for i, c in enumerate(NEEDED))


class ExtractError(Exception):
    pass


# ---------------------------------------------------------------- utilidades de valores
def T(v):
    """Valor de celda -> texto recortado (igual que TrimStr del extractor local)."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    return str(v).strip()


def to_qty(v):
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except ValueError:
        return 0.0


def num(v):
    if v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return repr(float(v))


def q(s):
    """Cadena JSON ASCII-only; escapa < > & ' y todo lo no ASCII (pares sustitutos en minuscula)."""
    out = ['"']
    for ch in s:
        o = ord(ch)
        if ch == '"':
            out.append('\\"')
        elif ch == "\\":
            out.append("\\\\")
        elif ch == "\n":
            out.append("\\n")
        elif ch == "\r":
            out.append("\\r")
        elif ch == "\t":
            out.append("\\t")
        elif o < 0x20 or o > 0x7E or ch in "<>&'":
            if o > 0xFFFF:
                o -= 0x10000
                out.append("\\u%04x\\u%04x" % (0xD800 + (o >> 10), 0xDC00 + (o & 0x3FF)))
            else:
                out.append("\\u%04x" % o)
        else:
            out.append(ch)
    out.append('"')
    return "".join(out)


def pd(v):
    """PERIODO / DIA: numero si la celda es numerica, texto si viene como texto (.xlsb)."""
    if v is None:
        return "null"
    if isinstance(v, bool):
        return q(str(v))
    if isinstance(v, (int, float)):
        return num(v)
    return q(str(v))


def kv(name, val, first=False):
    return ("" if first else ",") + '"' + name + '":' + q(val)


# ---------------------------------------------------------------- lectura de libros
def check_headers(values):
    bad = []
    for i, exp in enumerate(EXPECTED_HEADERS):
        got = T(values[i]) if i < len(values) else ""
        if got.lower() != exp.lower():
            bad.append("col %d: '%s' (esperado '%s')" % (i + 1, got, exp))
    if bad:
        raise ExtractError("Encabezados de Base Stocks distintos al esquema:\n  " + "\n  ".join(bad))


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        raise ExtractError("Falta la libreria openpyxl (pip install -r tools/cloud/requirements.txt).")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if SHEET not in names:
            raise ExtractError("El archivo no tiene la hoja '%s'. Hojas: %s" % (SHEET, ", ".join(names)))
        ws = wb[SHEET]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        it = ws.iter_rows(min_row=1, values_only=True)
        try:
            header = next(it)
        except StopIteration:
            raise ExtractError("La hoja '%s' esta vacia." % SHEET)
        check_headers(list(header))
        rows = []
        for row in it:
            n = len(row)
            rows.append(tuple((row[c - 1] if c - 1 < n else None) for c in NEEDED))
    finally:
        wb.close()
    return names, rows


def read_xlsb(path):
    try:
        from pyxlsb import open_workbook
    except ImportError:
        raise ExtractError("Falta la libreria pyxlsb (pip install -r tools/cloud/requirements.txt).")
    with open_workbook(path) as wb:
        names = list(wb.sheets)
        if SHEET not in names:
            raise ExtractError("El archivo no tiene la hoja '%s'. Hojas: %s" % (SHEET, ", ".join(names)))
        header = {}
        rowmap = {}
        with wb.get_sheet(SHEET) as sh:
            for item in sh.rows(sparse=True):
                cells = [item] if hasattr(item, "c") else item
                for cell in cells:
                    if cell.r == 0:
                        header[cell.c] = cell.v
                        continue
                    idx = POS.get(cell.c + 1)
                    if idx is None:
                        continue
                    vals = rowmap.get(cell.r)
                    if vals is None:
                        vals = rowmap[cell.r] = [None] * len(NEEDED)
                    vals[idx] = cell.v
    if not header:
        raise ExtractError("La hoja '%s' esta vacia." % SHEET)
    check_headers([header.get(i) for i in range(39)])
    rows = [tuple(rowmap[r]) for r in sorted(rowmap)]
    return names, rows


def read_workbook(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsb":
        return read_xlsb(path)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    raise ExtractError("Formato no soportado (%s). Usa .xlsb o .xlsx." % ext)


# ---------------------------------------------------------------- logica de extraccion
def passes(row):
    """Predicado unico de alcance: Honor | Moviles+Accesorios | Nuevos | Disponible |
    PDV OPERATIVO | TIPO=Celular solo Moviles | USO Normal/Pack."""
    fam = T(row[POS[C_FAM]]).lower()
    is_mov = fam == "moviles"
    if not is_mov and fam != "accesorios":
        return False
    if T(row[POS[C_DISP]]).lower() != "disponible":
        return False
    if T(row[POS[C_ESTADOPDV]]).lower() != "operativo":
        return False
    if T(row[POS[C_MARCA]]).lower() != "honor":
        return False
    if T(row[POS[C_ESTADO]]).lower() != "nuevos":
        return False
    if is_mov and T(row[POS[C_TIPO]]).lower() != "celular":
        return False
    uso = T(row[POS[C_USO]]).lower()
    if uso != "normal" and uso != "pack":
        return False
    return True


def canon_map(rows, col):
    """Variante canonica por valor (sin distinguir mayusculas): gana la mas frecuente; en empate, la primera vista."""
    counts = {}
    for row in rows:
        v = T(row[POS[col]])
        if v:
            counts[v] = counts.get(v, 0) + 1
    canon = {}
    best = {}
    for k, cnt in counts.items():     # orden de insercion = orden de primera aparicion
        lk = k.lower()
        if lk not in best or cnt > best[lk]:
            best[lk] = cnt
            canon[lk] = k
    return canon


def norm(m, s):
    return m.get(s.lower(), s)


def build(rows, coverage=False):
    if not rows:
        raise ExtractError("La hoja Base Stocks no tiene filas de datos.")

    passing = [r for r in rows if passes(r)]

    map_fam = canon_map(passing, C_FAM)
    map_tipo = canon_map(passing, C_TIPO)
    map_sub = canon_map(passing, C_SUBTIPO)
    map_modelo = canon_map(passing, C_MODELO)
    map_prod = canon_map(passing, C_PROD)

    products = []            # producto, modelo, marca, familia, tipo, subtipo, catEq, catAcc, codOracle
    product_index = {}
    pdvs = []                # pdv, ubicacion, region, departamento, canal
    pdv_index = {}
    pair_index = {}
    pair_p, pair_s, pair_q = [], [], []
    kept = 0
    total_qty = 0.0

    for row in passing:
        pdv = T(row[POS[C_PDV]])
        if not pdv:
            continue
        producto = norm(map_prod, T(row[POS[C_PROD]]))
        pkey = producto.lower()
        p_idx = product_index.get(pkey)
        if p_idx is None:
            products.append((
                producto, norm(map_modelo, T(row[POS[C_MODELO]])), T(row[POS[C_MARCA]]),
                norm(map_fam, T(row[POS[C_FAM]])), norm(map_tipo, T(row[POS[C_TIPO]])),
                norm(map_sub, T(row[POS[C_SUBTIPO]])), T(row[POS[C_CATEQ]]), T(row[POS[C_CATACC]]),
                T(row[POS[C_CODORACLE]])))
            p_idx = len(products) - 1
            product_index[pkey] = p_idx

        skey = pdv.lower()
        s_idx = pdv_index.get(skey)
        if s_idx is None:
            pdvs.append((pdv, T(row[POS[C_UBIC]]), T(row[POS[C_REGION]]),
                         T(row[POS[C_DEPTO]]), T(row[POS[C_CANAL]])))
            s_idx = len(pdvs) - 1
            pdv_index[skey] = s_idx

        qty = to_qty(row[POS[C_ITEMS]])
        key = (p_idx, s_idx)
        pi = pair_index.get(key)
        if pi is None:
            pair_index[key] = len(pair_p)
            pair_p.append(p_idx)
            pair_s.append(s_idx)
            pair_q.append(qty)
        else:
            pair_q[pi] += qty
        total_qty += qty
        kept += 1

    out = ['{"periodo":', pd(rows[0][POS[C_PERIODO]]), ',"dia":', pd(rows[0][POS[C_DIA]])]

    out.append(',"products":[')
    for i, p in enumerate(products):
        if i:
            out.append(",")
        out.append("{" + kv("producto", p[0], True) + kv("modelo", p[1]) + kv("marca", p[2]) +
                   kv("familia", p[3]) + kv("tipo", p[4]) + kv("subtipo", p[5]) +
                   kv("catEq", p[6]) + kv("catAcc", p[7]) + kv("codOracle", p[8]) + "}")
    out.append('],"pdvs":[')
    for i, p in enumerate(pdvs):
        if i:
            out.append(",")
        out.append("{" + kv("pdv", p[0], True) + kv("ubicacion", p[1]) + kv("region", p[2]) +
                   kv("departamento", p[3]) + kv("canal", p[4]) + "}")
    out.append('],"stock":[')
    for i in range(len(pair_p)):
        if i:
            out.append(",")
        out.append("[%d,%d,%s]" % (pair_p[i], pair_s[i], num(pair_q[i])))
    out.append("]")

    stats = {"kept": kept, "products": len(products), "pdvs": len(pdvs), "pairs": len(pair_p),
             "total_qty": total_qty, "cov_models": 0, "cov_pdvs": 0}

    if coverage:
        # catalogo de modelos: Moviles primero, luego alfabetico (comparacion ordinal sin mayusculas)
        seen = set()
        models = []
        for p in products:
            mk = p[1].lower()
            if mk not in seen:
                seen.add(mk)
                models.append((p[1], p[3]))
        models.sort(key=lambda m: (0 if m[1] == "Moviles" else 1, m[0].upper()))
        model_index = dict((m[0].lower(), i) for i, m in enumerate(models))

        # universo: cualquier fila Honor con PDV (sin exigir disponibilidad/operatividad/estado/uso)
        univ_index = {}
        univ = []            # pdv, ubicacion, departamento, canal, estadoPdv
        for row in rows:
            if T(row[POS[C_MARCA]]).lower() != "honor":
                continue
            pdv = T(row[POS[C_PDV]])
            if not pdv:
                continue
            uk = pdv.lower()
            if uk in univ_index:
                continue
            univ.append((pdv, T(row[POS[C_UBIC]]), T(row[POS[C_DEPTO]]),
                         T(row[POS[C_CANAL]]), T(row[POS[C_ESTADOPDV]])))
            univ_index[uk] = len(univ) - 1
        have = [set() for _ in univ]
        for row in passing:
            pdv = T(row[POS[C_PDV]])
            if not pdv:
                continue
            ui = univ_index.get(pdv.lower())
            if ui is None:
                continue
            mi = model_index.get(norm(map_modelo, T(row[POS[C_MODELO]])).lower())
            if mi is None:
                continue
            have[ui].add(mi)

        out.append(',"coverage":{"models":[')
        for i, m in enumerate(models):
            if i:
                out.append(",")
            out.append("{" + kv("modelo", m[0], True) + kv("familia", m[1]) + "}")
        out.append('],"pdvs":[')
        for i, u in enumerate(univ):
            if i:
                out.append(",")
            out.append("{" + kv("pdv", u[0], True) + kv("ubicacion", u[1]) + kv("departamento", u[2]) +
                       kv("canal", u[3]) + kv("estadoPdv", u[4]) +
                       ',"have":[' + ",".join(str(x) for x in sorted(have[i])) + "]}")
        out.append("]}")
        stats["cov_models"] = len(models)
        stats["cov_pdvs"] = len(univ)

    out.append("}")
    return "".join(out), stats


def extract_text(path, coverage=False):
    names, rows = read_workbook(path)
    text, stats = build(rows, coverage)
    stats["sheets"] = names
    return text, stats


def main(argv=None):
    ap = argparse.ArgumentParser(description="Extrae stock_data.json desde la hoja Base Stocks (sin Excel).")
    ap.add_argument("--input", required=True, help="archivo .xlsb o .xlsx")
    ap.add_argument("--out", required=True, help="ruta del stock_data.json a escribir")
    ap.add_argument("--coverage", action="store_true", help="incluye el bloque 'coverage'")
    a = ap.parse_args(argv)
    try:
        if not os.path.isfile(a.input):
            raise ExtractError("No existe el archivo: %s" % a.input)
        text, s = extract_text(a.input, a.coverage)
    except ExtractError as e:
        sys.stderr.write("ERROR: %s\n" % e)
        return 1
    with open(a.out, "w", encoding="utf-8", newline="") as f:
        f.write(text)
    print("Hojas del libro: %s" % ", ".join(s["sheets"]))
    print("Filtered rows kept: %d" % s["kept"])
    print("Distinct products: %d" % s["products"])
    print("Distinct PDVs: %d" % s["pdvs"])
    print("Distinct stock pairs: %d" % s["pairs"])
    if a.coverage:
        print("Cobertura: modelos=%d universo PDV=%d" % (s["cov_models"], s["cov_pdvs"]))
    print("Wrote JSON to %s (%.1f KB)" % (a.out, os.path.getsize(a.out) / 1024.0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
